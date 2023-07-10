import datetime
from docx import Document
import openai

# Set up OpenAI API credentials
openai.api_key = 'insert OpenAI API key here'

def generate_personalized_line(company_name, position_title):
    prompt = f"Please write 1 simplified sentence to insert into my cover letter about why I'm interested in working for the company {company_name} as a {position_title}, based on the company's specific products, values, and mission. It should be specific to the company, personal and make me stand out as a candidate."
    response = openai.ChatCompletion.create(
        model="gpt-3.5-turbo",
        messages=[
            {"role": "user", "content": prompt}],
        max_tokens=100,
        temperature=0.5,
    )
    personalized_line = response.choices[0].message.content.strip()

    return personalized_line

def generate_cover_letter(company_name=None, position_title=None, template_file="template.docx"):
    if company_name is None:
        company_name = input("Enter the company name: ")
    if position_title is None:
        position_title = input("Enter the position title: ")

    today = datetime.date.today().strftime("%B %d, %Y")
    personalized_line = generate_personalized_line(company_name, position_title)

    # Load the template document
    template = Document(template_file)

    # Replace placeholders in the template
    for paragraph in template.paragraphs:
        for run in paragraph.runs:
            text = run.text.replace("<COMPANY_NAME>", company_name)
            text = text.replace("<POSITION_TITLE>", position_title)
            text = text.replace("<TODAY_DATE>", today)
            text = text.replace("<PERSONALIZED_LINE>", personalized_line)
            run.text = text

    # Generate output file name
    output_filename = f"FirstNameLastName_{company_name}_{position_title}_CoverLetter.docx"

    # Save the generated cover letter
    template.save(output_filename)
    print(f"Generated cover letter has been saved to {output_filename}")

# Example usage
generate_cover_letter()